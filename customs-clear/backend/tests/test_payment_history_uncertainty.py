"""Saved estimates must remain estimates in lists, batch summaries and downloads."""
from __future__ import annotations

import asyncio
import csv
import io
import json
from copy import deepcopy
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import CustomsCalculationHistory
from app.services import calculation_history_service as history
from app.services.export_service import generate_final_customs_excel
from app.services.payment_result_status import aggregate_payment_metadata, payment_result_metadata


WARNING = "Применимость преференции не подтверждена для товара и даты; сумма предварительная."


def payment(*, pending: bool = True) -> dict:
    result = {
        "status": "REVIEW_REQUIRED" if pending else "OK",
        "hs_code": "7112300000",
        "amounts_provisional": pending,
        "breakdown": {"total_payable": 152000.0},
        "data_quality": {"confidence": "exact", "amounts_provisional": pending},
    }
    if pending:
        result["tariff_preference"] = {
            "applied": False, "status": "needs_review", "reason": WARNING,
            "eligibility_verified": False, "candidate_duty_coefficient": 0.75,
        }
    return result


@pytest.fixture
def isolated_history(monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    CustomsCalculationHistory.__table__.create(engine)
    monkeypatch.setattr(history, "SessionLocal", sessionmaker(bind=engine))
    monkeypatch.setattr(history, "engine", engine)
    yield
    engine.dispose()


@pytest.mark.parametrize("pending", [True, False])
def test_saved_payment_list_detail_json_and_csv_preserve_recorded_state(isolated_history, pending):
    output = payment(pending=pending)
    original = deepcopy(output)
    calc_id = history.save_calculation_record(
        input_payload={"hs_code": "7112300000"}, output_payload=output,
    )
    # Later modification of the caller's dictionary must not rewrite saved evidence.
    output["amounts_provisional"] = not pending
    detail = history.get_calculation_record(calc_id)
    assert detail["output_payload"] == original
    listed = history.list_calculation_history()[0]
    exported = history.export_calculation_history_rows(full_json=True)[0]
    csv_row = next(csv.DictReader(io.StringIO(history.calculation_history_as_csv([exported]))))
    for row in (detail, listed, exported):
        assert row["amounts_provisional"] is pending
        assert row["payment_status"] == ("REVIEW_REQUIRED" if pending else "OK")
        assert row["total_payable"] == 152000
        assert row["tariff_preference_warning"] == (WARNING if pending else None)
    assert exported["output_payload"] == original
    assert csv_row["amounts_provisional"] == str(pending)
    assert csv_row["payment_status"] == original["status"]
    assert csv_row["tariff_preference_warning"] == (WARNING if pending else "")


@pytest.mark.parametrize("legacy", [
    {"breakdown": {"total_payable": 123.0}},
    {"status": "OK", "breakdown": {"total_payable": 123.0}},
    {"status": "OK", "breakdown": {"total_payable": 123.0}, "tariff_preference": {"applied": True, "duty_coefficient": 0.75}},
])
def test_older_history_keeps_unknown_finality_without_rewriting_legacy_json(isolated_history, legacy):
    calc_id = history.save_calculation_record(input_payload={}, output_payload=legacy)
    detail = history.get_calculation_record(calc_id)
    assert detail["output_payload"] == legacy
    assert detail["amounts_provisional"] is None
    assert detail["payment_status"] == legacy.get("status")
    assert detail["tariff_preference_warning"] is None
    assert history.list_calculation_history()[0]["amounts_provisional"] is None
    row = next(csv.DictReader(io.StringIO(history.calculation_history_as_csv(history.export_calculation_history_rows()))))
    assert row["amounts_provisional"] == ""


@pytest.mark.parametrize("kind,output", [
    ("compliance", {"status": "OK", "items_summary": [
        {"hs_code": "7112300000", "total_payable": 10, **payment_result_metadata(payment())},
        {"hs_code": "7112910000", "total_payable": 20, **payment_result_metadata(payment(pending=False))},
    ]}),
    ("copilot_batch", {"payments": [
        {"total": 10, **payment_result_metadata(payment())},
        {"total": 20, **payment_result_metadata(payment(pending=False))},
    ]}),
    ("compare", {"status": "OK", "scenarios": [{"profile": payment()}, {"profile": payment(pending=False)}]}),
])
def test_one_pending_item_keeps_entire_history_summary_under_review(isolated_history, kind, output):
    calc_id = history.save_calculation_record(input_payload={}, output_payload=output, kind=kind)
    row = history.get_calculation_record(calc_id)
    assert row["payment_status"] == "REVIEW_REQUIRED"
    assert row["amounts_provisional"] is True
    assert row["tariff_preference_warning"] == WARNING
    assert row["total_payable"] == (None if kind == "compare" else 30)


def test_batch_unknown_member_cannot_become_confirmed():
    result = aggregate_payment_metadata([payment(pending=False), {"total": 10}])
    assert result["amounts_provisional"] is None
    assert result["payment_status"] is None
    assert aggregate_payment_metadata([])["amounts_provisional"] is None


@pytest.mark.parametrize("absent", [None, "not a number", float("inf")])
def test_batch_missing_amount_is_not_added_as_zero(isolated_history, absent):
    calc_id = history.save_calculation_record(input_payload={}, kind="copilot_batch", output_payload={
        "payments": [{"total": 100, **payment_result_metadata(payment(pending=False))}, {"total": absent}],
    })
    detail = history.get_calculation_record(calc_id)
    assert detail["total_payable"] is None
    assert detail["amounts_provisional"] is None


def test_warning_can_survive_a_compact_profile_in_data_quality():
    result = payment_result_metadata({"status": "OK", "data_quality": {
        "amounts_provisional": True, "tariff_preference_warning": WARNING,
    }})
    assert result == {
        "payment_status": "REVIEW_REQUIRED", "amounts_provisional": True,
        "tariff_preference_warning": WARNING, "payment_review_reason": WARNING,
        "payment_review_reasons": [],
    }


def test_generic_missing_source_reason_survives_history_and_exports_without_becoming_preference(isolated_history):
    reason = "Источник ставки НДС не найден."
    output = {
        "status": "REVIEW_REQUIRED", "amounts_provisional": True,
        "breakdown": {"total_payable": 123},
        "payment_review_reason": reason, "payment_review_reasons": ["missing_vat_source"],
    }
    calc_id = history.save_calculation_record(input_payload={}, output_payload=output)
    row = history.get_calculation_record(calc_id)
    assert row["payment_review_reason"] == reason
    assert row["payment_review_reasons"] == ["missing_vat_source"]
    assert row["tariff_preference_warning"] is None
    exported = next(csv.DictReader(io.StringIO(history.calculation_history_as_csv([row]))))
    assert exported["payment_review_reason"] == reason
    assert json.loads(exported["payment_review_reasons"]) == ["missing_vat_source"]
    assert exported["tariff_preference_warning"] == ""
    ws = load_workbook(io.BytesIO(generate_final_customs_excel([{"payment_profile": output}])), data_only=True).active
    xlsx_row = dict(zip([cell.value for cell in ws[1]], [cell.value for cell in ws[2]]))
    assert xlsx_row["Причина проверки расчёта"] == reason
    assert xlsx_row["Причина проверки преференции"] is None
    assert xlsx_row["Итого платежей"] is None


@pytest.mark.parametrize("pending", [True, False, None])
def test_xlsx_separates_estimate_from_final_total_and_retains_warning(pending):
    profile = payment(pending=pending is True)
    if pending is None:
        profile.pop("amounts_provisional")
        profile["data_quality"].pop("amounts_provisional")
    raw = generate_final_customs_excel([{"payment_profile": profile}])
    ws = load_workbook(io.BytesIO(raw), data_only=True).active
    row = dict(zip([cell.value for cell in ws[1]], [cell.value for cell in ws[2]]))
    assert row["Итого платежей"] == (152000 if pending is False else None)
    assert row["Предварительная оценка платежей"] == (None if pending is False else 152000)
    assert row["Суммы предварительные"] == {True: "Да", False: "Нет", None: "Не сохранено"}[pending]
    assert row["Причина проверки преференции"] == (WARNING if pending is True else None)
    if pending is True:
        assert row["Статус расчёта"] == "REVIEW_REQUIRED"
    if pending is None:
        assert "не сохранён" in row["Статус расчёта"]


def test_xlsx_api_schema_does_not_discard_the_uncertainty_contract():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api.calculator import router

    app = FastAPI()
    app.include_router(router, prefix="/calculator")
    profile = payment()
    profile["breakdown"].update({
        "base_duty": 10000, "vat": 140000, "excise": 0, "anti_dumping": 0, "customs_fee": 2000,
    })
    response = TestClient(app).post("/calculator/export-excel", json={"items": [{"payment_profile": profile}]})
    assert response.status_code == 200
    ws = load_workbook(io.BytesIO(response.content), data_only=True).active
    row = dict(zip([cell.value for cell in ws[1]], [cell.value for cell in ws[2]]))
    assert row["Итого платежей"] is None
    assert row["Предварительная оценка платежей"] == 152000
    assert row["Статус расчёта"] == "REVIEW_REQUIRED"
    assert row["Причина проверки преференции"] == WARNING


def test_compliance_response_and_saved_summary_keep_payment_review(monkeypatch, isolated_history):
    from app.api import compliance

    monkeypatch.setattr(compliance, "compute_payments", lambda _: payment())
    monkeypatch.setattr(compliance, "check_position_non_tariff", AsyncMock(return_value={"status": "OK"}))
    req = compliance.ComplianceRequest(items=[{
        "hs_code": "7112300000", "description": "Товар", "country": "BR", "customs_value": 100000,
    }], save_history=True)
    response = asyncio.run(compliance.compliance_check(req))
    body = json.loads(response.body)
    assert body["status"] == "WARNING"  # No false clean result, no enforcement ERROR.
    assert body["meta"]["any_manual_review"] is True
    assert body["meta"]["amounts_provisional"] is True
    assert WARNING in body["items"][0]["risks"]
    row = history.list_calculation_history(kind="compliance")[0]
    detail = history.get_calculation_record(row["id"])
    summary = detail["output_payload"]["items_summary"][0]
    assert summary["tariff_preference"] == payment()["tariff_preference"]
    assert row["amounts_provisional"] is True
    assert row["tariff_preference_warning"] == WARNING


def test_copilot_batch_saved_summary_does_not_drop_payment_review(monkeypatch, isolated_history):
    from starlette.requests import Request
    from app.api import assistant

    monkeypatch.setattr(assistant, "run_copilot_batch", AsyncMock(return_value={
        "bundles": [{"effective_hs_code": "7112300000", "payment": payment()},
                    {"effective_hs_code": "7112910000", "payment": payment(pending=False)}],
        "merged_context_for_ai": {},
    }))
    monkeypatch.setattr(assistant, "analyze_copilot_bundle", AsyncMock(return_value={"status": "OK"}))
    monkeypatch.setattr(assistant, "rag_context_for_copilot", AsyncMock(return_value={}))
    monkeypatch.setattr(assistant, "similar_decisions_context", lambda *args, **kwargs: {})
    monkeypatch.setattr(assistant, "append_audit", lambda _: None)
    monkeypatch.setattr(assistant, "request_audit_meta", lambda _: {})
    req = assistant.CopilotBatchRequest(items=[{"hs_code": "7112300000", "description": "Товар"}],
        run_payment=True, save_calculation_history=True)
    request = Request({"type": "http", "headers": []})
    asyncio.run(assistant.assistant_copilot_batch(req, request))
    row = history.list_calculation_history(kind="copilot_batch")[0]
    assert row["amounts_provisional"] is True
    assert row["total_payable"] == 304000
    detail = history.get_calculation_record(row["id"])
    saved = detail["output_payload"]["payments"]
    assert saved[0]["tariff_preference"] == payment()["tariff_preference"]
    assert saved[0]["tariff_preference_warning"] == WARNING
    assert saved[1]["amounts_provisional"] is False
